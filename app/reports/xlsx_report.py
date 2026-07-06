"""
XLSX report generator using Pandas and OpenPyXL.
"""
from typing import List
import pandas as pd
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.models.health_check_result import HealthCheckResult, HealthStatus
from app.utils.logger import get_logger

logger = get_logger(__name__)


class XLSXReportGenerator:
    """Generates XLSX health check reports with formatting."""
    
    # Color schemes for different statuses
    STATUS_COLORS = {
        HealthStatus.HEALTHY.value: 'C6EFCE',      # Light green
        HealthStatus.WARNING.value: 'FFEB9C',      # Light yellow
        HealthStatus.CRITICAL.value: 'FFC7CE',     # Light red
        HealthStatus.FAILED.value: 'D9D9D9',       # Light gray
        HealthStatus.UNKNOWN.value: 'BDD7EE',      # Light blue
    }
    
    def generate(self, results: List[HealthCheckResult], output_file: str) -> None:
        """
        Generate XLSX report from health check results.
        
        Args:
            results: List of health check results
            output_file: Output file path
        """
        try:
            output_path = Path(output_file)
            output_path.parent.mkdir(parents=True, exist_ok=True)
            
            # Create Excel writer
            with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
                # Generate summary sheet
                self._create_summary_sheet(results, writer)
                
                # Generate detailed results sheet
                self._create_details_sheet(results, writer)
                
                # Generate services sheet
                self._create_services_sheet(results, writer)
                
                # Generate database sheet if applicable
                if any(r.database_status for r in results):
                    self._create_database_sheet(results, writer)
            
            # Apply formatting
            self._apply_formatting(output_path)
            
            logger.info(f"XLSX report saved to {output_file}")
            
        except Exception as e:
            logger.error(f"Error generating XLSX report: {e}")
            raise
    
    def _create_summary_sheet(self, results: List[HealthCheckResult], writer: pd.ExcelWriter) -> None:
        """Create summary sheet with overall statistics."""
        summary_data = {
            'Metric': [
                'Total Servers',
                'Healthy',
                'Warning',
                'Critical',
                'Failed',
                'Total Services Monitored',
                'Services Running',
                'Services Down'
            ],
            'Count': [
                len(results),
                sum(1 for r in results if r.overall_status == HealthStatus.HEALTHY),
                sum(1 for r in results if r.overall_status == HealthStatus.WARNING),
                sum(1 for r in results if r.overall_status == HealthStatus.CRITICAL),
                sum(1 for r in results if r.overall_status == HealthStatus.FAILED),
                sum(len(r.service_statuses) for r in results),
                sum(sum(1 for s in r.service_statuses if s.is_running) for r in results),
                sum(sum(1 for s in r.service_statuses if not s.is_running) for r in results),
            ]
        }
        
        df_summary = pd.DataFrame(summary_data)
        df_summary.to_excel(writer, sheet_name='Summary', index=False)
    
    def _create_details_sheet(self, results: List[HealthCheckResult], writer: pd.ExcelWriter) -> None:
        """Create detailed results sheet."""
        rows = []
        
        for result in results:
            row = {
                'Server ID': result.server_id,
                'Server Name': result.server_name,
                'IP Address': result.ip_address,
                'Hostname': result.hostname or 'N/A',
                'OS Type': result.os_type,
                'Environment': result.environment or 'N/A',
                'CPU Usage (%)': round(result.cpu_metric.usage, 2) if result.cpu_metric else None,
                'CPU Status': result.cpu_metric.status.value if result.cpu_metric else 'N/A',
                'Memory Usage (%)': round(result.memory_metric.usage, 2) if result.memory_metric else None,
                'Memory Status': result.memory_metric.status.value if result.memory_metric else 'N/A',
                'Disk Usage (%)': round(result.disk_metric.usage, 2) if result.disk_metric else None,
                'Disk Status': result.disk_metric.status.value if result.disk_metric else 'N/A',
                'Overall Status': result.overall_status.value,
                'Execution Time': result.execution_time.strftime('%Y-%m-%d %H:%M:%S'),
                'Duration (sec)': round(result.duration_seconds, 2) if result.duration_seconds else None,
                'Error': result.error_message or 'None',
            }
            rows.append(row)
        
        df_details = pd.DataFrame(rows)
        df_details.to_excel(writer, sheet_name='Server Details', index=False)
    
    def _create_services_sheet(self, results: List[HealthCheckResult], writer: pd.ExcelWriter) -> None:
        """Create services status sheet."""
        rows = []
        
        for result in results:
            for service in result.service_statuses:
                row = {
                    'Server Name': result.server_name,
                    'IP Address': result.ip_address,
                    'Service Name': service.service_name,
                    'Status': service.status.value,
                    'Running': 'Yes' if service.is_running else 'No',
                    'Message': service.message or 'N/A',
                }
                rows.append(row)
        
        if rows:
            df_services = pd.DataFrame(rows)
            df_services.to_excel(writer, sheet_name='Services', index=False)
    
    def _create_database_sheet(self, results: List[HealthCheckResult], writer: pd.ExcelWriter) -> None:
        """Create database connectivity sheet."""
        rows = []
        
        for result in results:
            if result.database_status:
                row = {
                    'Server Name': result.server_name,
                    'IP Address': result.ip_address,
                    'Database Type': result.database_status.database_type,
                    'Database Host': result.database_status.host,
                    'Database Port': result.database_status.port,
                    'Connected': 'Yes' if result.database_status.is_connected else 'No',
                    'Status': result.database_status.status.value,
                    'Response Time (ms)': round(result.database_status.response_time_ms, 2) if result.database_status.response_time_ms else None,
                    'Error': result.database_status.error_message or 'None',
                }
                rows.append(row)
        
        if rows:
            df_database = pd.DataFrame(rows)
            df_database.to_excel(writer, sheet_name='Database', index=False)
    
    def _apply_formatting(self, file_path: Path) -> None:
        """Apply conditional formatting and styling to Excel file."""
        try:
            workbook = load_workbook(file_path)
            
            # Format each sheet
            for sheet_name in workbook.sheetnames:
                sheet = workbook[sheet_name]
                
                # Apply header formatting
                self._format_header(sheet)
                
                # Apply conditional formatting for status columns
                self._apply_conditional_formatting(sheet)
                
                # Auto-adjust column widths
                self._auto_adjust_columns(sheet)
            
            workbook.save(file_path)
            
        except Exception as e:
            logger.warning(f"Could not apply formatting: {e}")
    
    def _format_header(self, sheet) -> None:
        """Format header row."""
        header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF')
        alignment = Alignment(horizontal='center', vertical='center')
        
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = alignment
    
    def _apply_conditional_formatting(self, sheet) -> None:
        """Apply conditional formatting based on status values."""
        # Find status columns
        status_columns = []
        for col_idx, cell in enumerate(sheet[1], 1):
            if cell.value and 'Status' in str(cell.value):
                status_columns.append(col_idx)
        
        # Apply formatting to status columns
        for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row):
            for col_idx in status_columns:
                cell = row[col_idx - 1]
                status_value = str(cell.value)
                
                if status_value in self.STATUS_COLORS:
                    fill_color = self.STATUS_COLORS[status_value]
                    cell.fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type='solid')
    
    def _auto_adjust_columns(self, sheet) -> None:
        """Auto-adjust column widths based on content."""
        for column in sheet.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            
            for cell in column:
                try:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                except:
                    pass
            
            adjusted_width = min(max_length + 2, 50)  # Cap at 50
            sheet.column_dimensions[column_letter].width = adjusted_width
